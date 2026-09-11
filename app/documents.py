import csv
import hashlib
import io
import re
from collections import deque
from dataclasses import dataclass
from pathlib import PurePosixPath
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import settings
from app.errors import DocumentNotFoundError, InsufficientEvidenceError, UnsupportedDocumentError
from app.retrieval import hybrid_rank, tokenize
from app.schemas import (
    DocumentCitation,
    DocumentRecord,
    DocumentResearchRequest,
    DocumentResearchResult,
    DocumentType,
    ExtractedField,
)

SUPPORTED_SUFFIXES = {".pdf", ".csv", ".txt", ".md", ".xlsx"}
_VALUE = r"(?:R\$\s*)?[-+]?(?:\d{1,3}(?:[.\s]\d{3})+|\d+)(?:[,.]\d+)?(?:\s*(?:milh(?:ão|ões)|bilh(?:ão|ões)))?"
_FIELD_PATTERNS = {
    "revenue": rf"(?:receita\s+(?:líquida|bruta)|revenue)\s*[:\-]?\s*({_VALUE})",
    "net_income": rf"(?:lucro\s+líquido|net\s+income)\s*[:\-]?\s*({_VALUE})",
    "total_debt": rf"(?:dívida\s+total|total\s+debt)\s*[:\-]?\s*({_VALUE})",
    "equity": rf"(?:patrimônio\s+líquido|equity)\s*[:\-]?\s*({_VALUE})",
    "earnings_per_share": rf"(?:lucro\s+por\s+ação|earnings\s+per\s+share|eps)\s*[:\-]?\s*({_VALUE})",
}


@dataclass
class _Chunk:
    location: str
    text: str


@dataclass
class _StoredDocument:
    record: DocumentRecord
    chunks: list[_Chunk]


class InMemoryDocumentStore:
    """Bounded development store with a production-like hybrid retrieval contract."""

    def __init__(self, max_documents: int = 100) -> None:
        self._documents: deque[_StoredDocument] = deque(maxlen=max_documents)

    def save(self, document: _StoredDocument) -> DocumentRecord:
        self._documents.append(document)
        return document.record

    def list(self) -> list[DocumentRecord]:
        return [item.record for item in reversed(self._documents)]

    def get(self, document_id: str) -> _StoredDocument:
        for document in self._documents:
            if document.record.id == document_id:
                return document
        raise DocumentNotFoundError(f"document {document_id} was not found")

    def search(self, request: DocumentResearchRequest) -> DocumentResearchResult:
        candidates = list(self._documents)
        if request.document_id:
            candidates = [self.get(request.document_id)]
        if request.ticker:
            candidates = [item for item in candidates if item.record.ticker == request.ticker]

        keyed_chunks: dict[str, tuple[_StoredDocument, _Chunk]] = {}
        searchable: list[tuple[str, str]] = []
        for document in candidates:
            for index, chunk in enumerate(document.chunks):
                key = f"{document.record.id}:{index}"
                keyed_chunks[key] = (document, chunk)
                searchable.append((key, chunk.text))

        ranked = hybrid_rank(
            request.question,
            searchable,
            top_k=settings.retrieval_top_k,
            min_score=settings.retrieval_min_score,
        )
        if not ranked:
            raise InsufficientEvidenceError("no document evidence supports this question")

        terms = set(tokenize(request.question))
        citations: list[DocumentCitation] = []
        for match in ranked[:3]:
            document, chunk = keyed_chunks[match.key]
            citations.append(
                DocumentCitation(
                    document_id=document.record.id,
                    document_name=document.record.name,
                    location=chunk.location,
                    excerpt=_excerpt(chunk.text, terms),
                    score=match.combined_score,
                )
            )
        sources = ", ".join(f"{citation.document_name} ({citation.location})" for citation in citations)
        confidence = min(0.97, 0.45 + 0.5 * ranked[0].combined_score)
        return DocumentResearchResult(
            answer=f"Foram encontradas evidências documentais relevantes em {sources}.",
            confidence=confidence,
            citations=citations,
            retrieval_method="hybrid_bm25_hash_embedding_mmr",
        )


def ingest_document(
    *,
    name: str,
    content_type: str,
    data: bytes,
    ticker: str | None,
) -> _StoredDocument:
    suffix = PurePosixPath(name).suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedDocumentError(
            "supported document types are PDF, CSV, TXT, Markdown and XLSX"
        )
    pages = _extract_pages(suffix, data)
    non_empty_pages = [(location, text.strip()) for location, text in pages if text.strip()]
    if not non_empty_pages:
        raise UnsupportedDocumentError("the document has no extractable text")

    document_id = str(uuid4())
    chunks = _chunk_pages(non_empty_pages)
    document_type = _classify_document(suffix, "\n".join(text for _, text in non_empty_pages))
    fields = _extract_fields(document_id, name, non_empty_pages)
    confidence = 0.9 if fields else 0.7
    record = DocumentRecord(
        id=document_id,
        name=name,
        content_type=content_type or "application/octet-stream",
        ticker=ticker,
        document_type=document_type,
        page_count=len(non_empty_pages),
        chunk_count=len(chunks),
        extraction_confidence=confidence,
        extracted_fields=fields,
        source_sha256=hashlib.sha256(data).hexdigest(),
    )
    return _StoredDocument(record=record, chunks=chunks)


def _extract_pages(suffix: str, data: bytes) -> list[tuple[str, str]]:
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return [(f"page {index}", page.extract_text() or "") for index, page in enumerate(reader.pages, 1)]
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        pages: list[tuple[str, str]] = []
        for worksheet in workbook.worksheets:
            rows = [
                " | ".join(str(value) for value in row if value is not None)
                for row in worksheet.iter_rows(values_only=True)
            ]
            pages.append((f"sheet {worksheet.title}", "\n".join(row for row in rows if row)))
        return pages
    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        rows = csv.reader(io.StringIO(text))
        return [("row data", "\n".join(" | ".join(row) for row in rows))]
    return [("text", data.decode("utf-8", errors="replace"))]


def _chunk_pages(pages: list[tuple[str, str]], size: int = 1000, overlap: int = 160) -> list[_Chunk]:
    chunks: list[_Chunk] = []
    for location, text in pages:
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            chunks.append(_Chunk(location=location, text=text[start:end]))
            if end == len(text):
                break
            start = end - overlap
    return chunks


def _classify_document(suffix: str, text: str) -> DocumentType:
    normalized = text.casefold()
    if suffix in {".csv", ".xlsx"}:
        return DocumentType.SPREADSHEET
    if any(term in normalized for term in ("balanço patrimonial", "balance sheet", "receita líquida")):
        return DocumentType.FINANCIAL_STATEMENT
    if any(term in normalized for term in ("relatório", "report", "resultado")):
        return DocumentType.FINANCIAL_REPORT
    return DocumentType.UNKNOWN


def _extract_fields(
    document_id: str, document_name: str, pages: list[tuple[str, str]]
) -> list[ExtractedField]:
    extracted: list[ExtractedField] = []
    for field_name, pattern in _FIELD_PATTERNS.items():
        for location, text in pages:
            match = re.search(pattern, text, flags=re.IGNORECASE)
            if match:
                extracted.append(
                    ExtractedField(
                        name=field_name,
                        value=_parse_number(match.group(1)),
                        citation=DocumentCitation(
                            document_id=document_id,
                            document_name=document_name,
                            location=location,
                            excerpt=match.group(0),
                        ),
                    )
                )
                break
    return extracted


def _parse_number(value: str) -> float:
    normalized = value.casefold().replace("r$", "").strip()
    multiplier = 1
    if "bilh" in normalized:
        multiplier = 1_000_000_000
    elif "milh" in normalized:
        multiplier = 1_000_000
    normalized = re.sub(r"\s*(milh(?:ão|ões)|bilh(?:ão|ões))", "", normalized).strip()
    if "," in normalized:
        normalized = normalized.replace(".", "").replace(",", ".")
    else:
        normalized = normalized.replace(" ", "")
    return float(normalized) * multiplier


def _excerpt(text: str, terms: set[str], length: int = 280) -> str:
    normalized = text.casefold()
    positions = [normalized.find(term) for term in terms if normalized.find(term) >= 0]
    start = max(0, min(positions) - 80) if positions else 0
    return " ".join(text[start : start + length].split())
